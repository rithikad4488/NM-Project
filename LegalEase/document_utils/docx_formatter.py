from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt

from document_utils.common import logo_path, parse_terms, sanitize_text


def format_docx(
    text: str,
    doc_type: str,
    terms: str = "",
) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    logo = logo_path()
    if logo:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(logo), width=Inches(1.25))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(sanitize_text(doc_type).upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for block in sanitize_text(text).split("\n\n"):
        block = block.strip()
        if not block:
            continue

        lines = block.splitlines()
        if len(lines) == 1 and (
            lines[0].isupper() or lines[0][:3].isdigit()
        ):
            p = doc.add_paragraph()
            r = p.add_run(lines[0])
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        for index, line in enumerate(lines):
            if index:
                p.add_run().add_break()
            p.add_run(line)

    parsed_terms = parse_terms(terms)
    if parsed_terms:
        doc.add_paragraph()
        heading = doc.add_paragraph()
        r = heading.add_run("KEY TERMS")
        r.bold = True
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "No."
        hdr[1].text = "Term"
        for index, term in enumerate(parsed_terms, start=1):
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = term
            for cell in cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-assisted draft | Review with a qualified legal professional").italic = True

    output = BytesIO()
    doc.save(output)
    return output.getvalue()
